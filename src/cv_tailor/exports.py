"""Build downloadable DOCX and PDF versions of the tailored CV.

The export carries the honest coverage summary - hard requirements met,
gaps, and the provenance of the tailoring - never a fake ATS percentage.
"""
from __future__ import annotations

import io


def _coverage_lines(hard_covered: int, hard_total: int, preferred_covered: int, preferred_total: int,
                    matched: list[str], missing: list[str]) -> list[str]:
    lines = [f"Requirement coverage: {hard_covered} of {hard_total} hard requirements fully covered."]
    if preferred_total:
        lines.append(f"Preferred requirements: {preferred_covered} of {preferred_total} fully covered.")
    if matched:
        lines.append("Covered skills: " + ", ".join(matched))
    if missing:
        lines.append("Gaps (absent from the base CV, not added): " + ", ".join(missing))
    return lines


NOTE = (
    "Gaps are reported, not inserted. Every tailoring change links back to base-CV evidence, "
    "and every generated line passed the number, entity, and entailment gates."
)


def build_docx(tailored_cv: str, hard_covered: int, hard_total: int,
               preferred_covered: int, preferred_total: int,
               matched: list[str], missing: list[str]) -> bytes:
    import docx
    from docx.shared import Pt

    document = docx.Document()
    document.add_heading("Tailored CV", level=0)
    meta = document.add_paragraph()
    meta.add_run(f"Requirement coverage: {hard_covered}/{hard_total} hard requirements").italic = True
    for block in tailored_cv.strip().split("\n"):
        line = block.strip()
        if not line:
            continue
        if line.isupper() and len(line) < 60:
            document.add_heading(line.title(), level=2)
        else:
            document.add_paragraph(line)

    document.add_page_break()
    document.add_heading("Coverage report", level=1)
    for line in _coverage_lines(hard_covered, hard_total, preferred_covered, preferred_total, matched, missing):
        document.add_paragraph(line, style="List Bullet")
    note = document.add_paragraph()
    note.add_run(NOTE).font.size = Pt(9)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_pdf(tailored_cv: str, hard_covered: int, hard_total: int,
              preferred_covered: int, preferred_total: int,
              matched: list[str], missing: list[str]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    buffer = io.BytesIO()
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4, leftMargin=56, rightMargin=56, topMargin=56, bottomMargin=56
    )
    story = [Paragraph("Tailored CV", styles["Title"])]
    story.append(Paragraph(
        f"<i>Requirement coverage: {hard_covered}/{hard_total} hard requirements</i>", styles["Normal"]))
    story.append(Spacer(1, 14))

    for block in tailored_cv.strip().split("\n"):
        line = block.strip()
        if not line:
            story.append(Spacer(1, 6))
            continue
        style = styles["Heading2"] if line.isupper() and len(line) < 60 else styles["Normal"]
        story.append(Paragraph(_escape(line), style))

    story.append(Spacer(1, 18))
    story.append(Paragraph("Coverage report", styles["Heading1"]))
    for line in _coverage_lines(hard_covered, hard_total, preferred_covered, preferred_total, matched, missing):
        story.append(Paragraph(_escape(line), styles["Normal"]))
    story.append(Spacer(1, 8))
    story.append(Paragraph(f"<i>{_escape(NOTE)}</i>", styles["Normal"]))

    doc.build(story)
    return buffer.getvalue()


def _escape(text: str) -> str:
    return (
        text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    )
